"""Generate fictional sample resumes in resumes/ (PDF, DOCX, TXT).

Run:  python scripts/make_sample_resumes.py
Needs: reportlab, python-docx (both in requirements.txt)
"""

from pathlib import Path

import docx
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent.parent / "resumes"

RESUMES = {
    "resume_john_doe.pdf": [
        "John Doe",
        "Machine Learning Engineer | john.doe@example.com | Bengaluru, India",
        "",
        "SUMMARY",
        "ML engineer with 5 years of Python experience building recommendation",
        "and NLP systems in production.",
        "",
        "SKILLS",
        "Python, PyTorch, scikit-learn, SQL, Docker, AWS SageMaker, FastAPI",
        "",
        "EXPERIENCE",
        "Senior ML Engineer, ShopSmart (2022 - present)",
        "- Built a Python recommendation service serving 3M users/day.",
        "- Cut model inference latency 40% with ONNX and batching.",
        "ML Engineer, DataNest (2019 - 2022)",
        "- Developed NLP pipelines for support-ticket classification.",
        "",
        "EDUCATION",
        "B.Tech Computer Science, IIT Delhi (2019)",
    ],
    "resume_alex_chen.pdf": [
        "Alex Chen",
        "DevOps Engineer | alex.chen@example.com | Singapore",
        "",
        "SKILLS",
        "Kubernetes, Terraform, Go, Bash, AWS, GitHub Actions, Prometheus",
        "",
        "EXPERIENCE",
        "DevOps Engineer, CloudPeak (2020 - present)",
        "- Migrated 120 services to Kubernetes with zero downtime.",
        "- Wrote Go tooling for automated certificate rotation.",
        "",
        "EDUCATION",
        "B.Eng Computer Engineering, NUS (2020)",
    ],
}

DOCX_RESUME = ("resume_jane_smith.docx", [
    ("Jane Smith", "Title"),
    ("Data Engineer | jane.smith@example.com | London, UK", None),
    ("Skills", "Heading 1"),
    ("Python, Apache Spark, Airflow, dbt, Snowflake, SQL", None),
    ("Experience", "Heading 1"),
    ("Data Engineer, FinStream (2021 - present): built Python and Spark ETL pipelines "
     "processing 2 TB/day; introduced dbt testing that reduced data incidents by 60%.", None),
    ("Analyst, RetailCo (2018 - 2021): automated weekly reporting with SQL and Excel.", None),
    ("Education", "Heading 1"),
    ("MSc Data Science, University of Edinburgh (2018)", None),
])

TXT_RESUME = ("resume_priya_sharma.txt", """Priya Sharma
Frontend Developer | priya.sharma@example.com | Pune, India

SKILLS
JavaScript, TypeScript, React, Next.js, Tailwind CSS, Jest, Figma

EXPERIENCE
Frontend Developer, PixelWorks (2021 - present)
- Led the React rewrite of the customer dashboard (Lighthouse score 58 -> 94).
- Built a shared component library used by 6 product teams.

Junior Web Developer, StartHub (2019 - 2021)
- Built marketing sites with Next.js and headless CMS.

EDUCATION
B.E. Information Technology, Pune University (2019)
""")


def write_pdf(path: Path, lines: list[str]) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    y = 800
    for line in lines:
        is_heading = line.isupper() and line
        c.setFont("Helvetica-Bold" if is_heading or line == lines[0] else "Helvetica",
                  16 if line == lines[0] else 11)
        c.drawString(60, y, line)
        y -= 18
    c.save()


def main() -> None:
    OUT.mkdir(exist_ok=True)
    for name, lines in RESUMES.items():
        write_pdf(OUT / name, lines)

    name, paras = DOCX_RESUME
    d = docx.Document()
    for text, style in paras:
        d.add_paragraph(text, style=style) if style else d.add_paragraph(text)
    d.save(str(OUT / name))

    name, text = TXT_RESUME
    (OUT / name).write_text(text, encoding="utf-8")
    print(f"Wrote sample resumes to {OUT}")


if __name__ == "__main__":
    main()
